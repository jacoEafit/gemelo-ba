#!/usr/bin/env python3
"""Genera fuentes_gemelo_buenos_aires.xlsx: una fila por fuente única, desde lago/*.json (requiere openpyxl)."""
import glob, json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

PENDIENTES = {  # campos que existen o se necesitan y aún NO están en el lago
    "ba_barrios": "Población 2022 por barrio (no publicada en datos abiertos). Geometría en lago/geo/barrios.geojson",
    "ba_comunas": "Densidad 2022 (cruzar con censo). Geometría en lago/geo/comunas.geojson",
    "ba_data_ckan": "Ninguno: catálogo completo (455 conjuntos); falta monitorear cambios de licencia/actualización",
    "ba_demografia": "Edad y sexo por barrio, hogares, esperanza de vida (recursos del paquete sin curar); barrios aún Censo 2010",
    "dgeyc_censo2022": "Solo población por comuna. Faltan viviendas, hogares y servicios por comuna (informe 1877 los trae en gráficos) y radios censales",
    "ba_espacios_culturales": "Capacidad total y cantidad de salas por barrio (campos existen, no agregados); coordenadas no guardadas",
    "ba_obras_registradas": "Tipo de obra (acronimo/descripcion) por barrio; son permisos: las alturas vienen de ba_tejido_urbano",
    "ba_presupuesto": "Solo el último trimestre publicado (acumulado); falta ajuste por inflación y desglose por objeto del gasto",
    "ba_presupuesto_historico": "Solo el último año cerrado; 2013-2024 sin cargar (serie histórica en pesos corrientes)",
    "eoh_caba": "Tarifa media, estadía media y ocupación por categoría; se publica con ~1 año de rezago",
    "ba_espacios_verdes": "Arbolado; espacios verdes privados; m² por habitante (población 2022 por barrio no disponible)",
    "ba_establecimientos_educativos": "Matrícula (alumnos) y cargos docentes; el padrón trae correos y no se usa",
    "ba_hospitales": "Camas, egresos y guardias (otros conjuntos: movimiento hospitalario); CESAC y privados no incluidos",
    "ipcba": "Divisiones del IPCBA (alimentos, vivienda, transporte…); serie desde 2022-02 (base 2021)",
    "sipa_caba": "Por rama de actividad y salario medio; empleo público y no registrado no incluidos",
    "open_meteo": "Viento, humedad y radiación; dato de grilla (no estación, no distingue barrios)",
    "ba_cesac": "Especialidades por CESAC (campo especialid sin agregar); atenciones; teléfono/web no se copian",
    "ba_ecobici_estaciones": "Viajes en EcoBici (conjunto bicicletas-publicas, pesado) y capacidad de anclajes",
    "ba_colectivos_paradas": "Recorridos y frecuencias (GTFS); pasajeros por línea",
    "ba_puntos_verdes": "Materiales recibidos por punto (pesajes 2015-2019); el listado solo trae 17 puntos con atención",
    "ba_tejido_urbano": "Superficie de cada volumen (para m³ construidos) y geometrías 3D (~1 GB: GeoParquet/PMTiles, ver contrato); datos hasta 2021",
    "ba_parcelas": "Solo se usa smp → barrio; geometría de parcelas (329 MB) no cargada",
    "ba_subte_molinetes": "Franja horaria (cada 15 min) y años anteriores sin cargar; se guarda solo el agregado por estación y mes",
    "ba_subte_estaciones": "Accesibilidad; campo línea (A, B, C…) en subte_lineas.geojson. Viajes por estación: ver ba_subte_molinetes",
    "wikimedia_pageviews": "Otros artículos/idiomas; mide interés, no opinión ni sentimiento",
    "ba_delitos": "Tasa por 1.000 hab, franja horaria, subtipo, uso de arma; años 2016-2023 sin cargar",
    "ba_alojamientos": "Plazas/habitaciones (no están en este dataset); la demanda viene de la EOH (eoh_caba)",
    "ba_calidad_aire": "Estación Palermo sin lecturas válidas en 2025; año 2026; PM2.5 no se mide; índice de calidad",
    "ba_sensores_transito": "Lecturas en vivo (API de Transporte, suspendida por el portal); aquí solo inventario de 52 sensores",
    "ba_tmi": "Lecturas meteorológicas de las torres (temperatura, humedad); aquí solo inventario de 16 torres",
}
SIN_FUENTE = [  # temas del gemelo sin fuente curada: (tema, nombre, url, licencia, estado, qué falta)
    ("(demografía 2022 por barrio)", "Información Censal por Radio (BA Data)", "https://data.buenosaires.gob.ar/km/dataset/informacion-censal-por-radio",
     "CC-BY-2.5-AR", "insuficiente",
     "Solo radios de 2001 y 2010: no da población 2022 por barrio. Los radios 2022 no están en BA Data"),
    ("(reclamos)", "BA Data reclamos 147", "", "", "descartado", "Descartado por descarga pesada; faltan todos los campos"),
    ("(sensores en tiempo real)", "API de transporte/sensores", "", "", "pendiente", "El portal suspendió las API y GTFS (aviso del 2026-10-06); faltan lecturas"),
]
fuentes, usos = {}, {}
for ruta in sorted(glob.glob("lago/*.json")):
    d = json.load(open(ruta, encoding="utf-8"))
    for f in d["fuentes"]:
        fuentes.setdefault(f["id"], f)
        u = usos.setdefault(f["id"], {"temas": [], "cifras": [], "vig": set(), "probado": ""})
        u["temas"].append(d["tema"])
        u["probado"] = max(u["probado"], d["probado"])
        for k, c in d.get("cifras", {}).items():
            if c["fuente"] == f["id"]:
                u["cifras"].append(f"{k}={c['valor']:,} {c['unidad']}".replace(",", "."))
                u["vig"].add(str(c["vigencia"]))
        for k, s in {**d.get("series", {}), **d.get("desgloses", {})}.items():
            if s.get("fuente") == f["id"]:
                u["vig"].add(str(s["vigencia"] if "vigencia" in s else ""))

wb = Workbook()
ws = wb.active
ws.title = "fuentes"
ws.append(["#", "id", "nombre", "temas del lago", "url", "licencia", "estado", "probado", "vigencia", "cifras curadas",
           "campos pendientes"])
for i, (id_, f) in enumerate(sorted(fuentes.items()), 1):
    u = usos[id_]
    ws.append([i, id_, f["nombre"], ", ".join(dict.fromkeys(u["temas"])), f["url"], f["licencia"], f["estado"], u["probado"],
               ", ".join(sorted(v for v in u["vig"] if v)), "; ".join(dict.fromkeys(u["cifras"])), PENDIENTES.get(id_, "")])
n = len(fuentes)
for j, (tema, nombre, url, lic, estado, falta) in enumerate(SIN_FUENTE, n + 1):
    ws.append([j, "", nombre, tema, url, lic, estado, "", "", "", falta])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="1F4E78")
for col, w in zip("ABCDEFGHIJK", (4, 24, 36, 28, 55, 22, 11, 12, 14, 55, 80)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical="top")
ws.freeze_panes = "A2"
wb.save("fuentes_gemelo_buenos_aires.xlsx")
print(f"{n} fuentes vivas + {len(SIN_FUENTE)} sin fuente = {ws.max_row - 1} filas")
