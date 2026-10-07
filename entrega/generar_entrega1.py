#!/usr/bin/env python3
"""Entrega 1 · llena la plantilla del profesor (Diccionario_de_datos_TallerDatos.xlsx) desde el lago real.
Uso, desde la raíz del proyecto:  python entrega/generar_entrega1.py   (requiere openpyxl)
- Amplía las tres hojas de 40 a FILAS filas útiles (formato, listas desplegables y fórmulas del Resumen).
- Hoja 1: una fila por fuente del rastreo (integradas, candidatas, caídas y descartadas).
- Hoja 2: una fila por campo del lago; los campos que solo cambian por año o por línea van en una fila.
- Hoja 3: bitácora de la búsqueda con IA, con sus errores.
Ejemplos y unidades salen de lago/*.json: si el lago cambia, se vuelve a correr."""
import copy, glob, json, pathlib, re
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

RAIZ = pathlib.Path(__file__).resolve().parent.parent
PLANTILLA = RAIZ / "entrega" / "plantilla_Diccionario_de_datos_TallerDatos.xlsx"
SALIDA = RAIZ / "entrega" / "Diccionario_de_datos_Buenos_Aires.xlsx"
FILAS = 200  # última fila útil (la plantilla trae hasta la 44)
PRIMERA = 5  # la fila 4 es el ejemplo gris de Medellín: se deja

lago = {j["tema"]: j for j in (json.loads(pathlib.Path(p).read_text(encoding="utf-8")) for p in sorted(glob.glob(str(RAIZ / "lago/*.json"))))}
geo = {pathlib.Path(p).stem: json.loads(pathlib.Path(p).read_text(encoding="utf-8")) for p in sorted(glob.glob(str(RAIZ / "lago/geo/*.geojson")))}
fuente_lago = {f["id"]: f for j in lago.values() for f in j["fuentes"]}
probado_de = {}
for j in lago.values():
    for f in j["fuentes"]:
        probado_de[f["id"]] = max(probado_de.get(f["id"], ""), j["probado"])

CDN = "https://cdn.buenosaires.gob.ar/datosabiertos/datasets/"
SERIES_API = "https://apis.datos.gob.ar/series/api/series/?ids="
IA = "IA (Claude Code): búsqueda en el catálogo CKAN de BA Data; probada con descarga del script"
IA_API = "IA (Claude Code): búsqueda en la API de series de datos.gob.ar; probada con la consulta del script"
SONDEO = "IA (Claude Code): búsqueda en el catálogo CKAN de BA Data; probada leyendo la ficha del portal y el tamaño del archivo, sin descargarlo"
SONDEO_API = "IA (Claude Code): búsqueda en la API de series de datos.gob.ar; probada con la búsqueda de la API, sin descargar la serie"

# ---------------------------------------------------------------- 1. Fuentes integradas (orden = ID F01, F02…)
# id lago: (entidad, tipo publicador, URL exacta, formato, cómo se obtiene, licencia tal como aparece, personas,
#           cobertura, periodo, frecuencia, cómo la encontraron, observaciones)
INTEGRADAS = {
    "ba_barrios": ("GCBA · Secretaría de Innovación y Transformación Digital", "Entidad pública",
                   CDN + "innovacion-transformacion-digital/barrios/barrios.geojson", "GeoJSON", "Descarga directa", None, "No",
                   "Ciudad, 48 barrios", "Límites vigentes", "No se sabe", IA, "Polígonos simplificados a ~10 m en lago/geo/barrios.geojson"),
    "ba_comunas": ("GCBA · Secretaría de Innovación y Transformación Digital", "Entidad pública",
                   CDN + "innovacion-transformacion-digital/comunas/comunas.geojson", "GeoJSON", "Descarga directa", None, "No",
                   "Ciudad, 15 comunas", "Límites vigentes", "No se sabe", IA, "Polígonos simplificados en lago/geo/comunas.geojson"),
    "ba_demografia": ("GCBA · Dirección General de Estadísticas y Censos", "Entidad pública",
                      CDN + "barrios/caba_pob_barrios_2010.csv", "CSV",
                      "Descarga directa", None, "Solo conteos por zona", "Ciudad, 48 barrios", "Censo 2010",
                      "Única vez", IA, "Es la única población por barrio abierta: Censo 2010. No hay 2022 por barrio."),
    "dgeyc_censo2022": ("GCBA · DGEyC, con base en INDEC (informe 1877)", "Entidad pública",
                        "https://www.estadisticaciudad.gob.ar/eyc/wp-content/uploads/2024/07/ir_2024_1877.pdf", "PDF",
                        "Solo PDF o página", "no declara", "Solo conteos por zona", "Ciudad, 15 comunas", "Censo 2022 (definitivo)",
                        "Única vez", "IA (Claude Code): búsqueda web del informe oficial; cifras transcritas a mano y verificadas contra el total",
                        "Estado «declarado»: Anexo 1, cuadro C.1, transcrito en ingesta/datos/censo2022_comunas.csv; la suma de comunas = total CABA"),
    "ba_delitos": ("GCBA · Ministerio de Justicia y Seguridad", "Entidad pública",
                   CDN + "ministerio-de-justicia-y-seguridad/delitos/delitos_2024.csv ; " + CDN + "ministerio-de-justicia-y-seguridad/delitos/delitos_2025.csv",
                   "CSV", "Descarga directa", None, "Personas identificables", "Ciudad, por barrio y comuna", "2024-01 a 2025-12",
                   "Anual", IA, "Cada fila es un hecho con fecha, franja y coordenadas: una víctima puede ser determinable. Al lago solo pasan conteos por barrio, tipo y mes. Hay archivos desde 2016 sin cargar."),
    "ba_presupuesto": ("GCBA · Ministerio de Hacienda y Finanzas", "Entidad pública",
                       CDN + "ministerio-de-economia-y-finanzas/presupuesto-ejecutado-2026/presupuesto-ejecutado-2026-2.csv", "CSV",
                       "Descarga directa", None, "No", "Ciudad (jurisdicción, finalidad, ubicación geográfica)", "2026-01 a 2026-06 (acumulado)",
                       "Trimestral", IA, "Montos con punto de miles; el trimestre es acumulado desde enero."),
    "ba_presupuesto_historico": ("GCBA · Ministerio de Hacienda y Finanzas", "Entidad pública",
                                 CDN + "ministerio-de-economia-y-finanzas/presupuesto-ejecutado/presupuesto-ejecutado-2025-4.csv", "CSV",
                                 "Descarga directa", None, "No", "Ciudad", "2025 (año cerrado)", "Trimestral", IA,
                                 "Mismo ministerio que la anterior, otro conjunto. Trae 2013-2025; solo se cargó 2025."),
    "ipcba": ("GCBA · Dirección General de Estadísticas y Censos (republicado por la Subsecretaría de Programación Macroeconómica)",
              "Entidad pública", SERIES_API + "193.2_NIVEL_GENERAL_2021_0_13_2", "JSON", "API sin llave",
              "Creative Commons Attribution 4.0", "No", "Ciudad", "2022-02 a 2026-08", "Mensual", IA_API,
              "Se prefirió al IPC-GBA del INDEC porque mide precios de la Ciudad. Se usa para deflactar el presupuesto."),
    "sipa_caba": ("Ministerio de Economía de la Nación (SIPA)", "Entidad pública", SERIES_API + "154.2_CABA_S_0_0_4", "JSON",
                  "API sin llave", "Creative Commons Attribution 4.0", "Solo conteos por zona", "Ciudad (lugar de trabajo)",
                  "2018-01 a 2026-06", "Mensual", IA_API, "Cuenta puestos con lugar de trabajo en CABA, no residentes."),
    "ba_alojamientos": ("GCBA · Ente de Turismo (ENTUR)", "Entidad pública",
                        CDN + "ente-de-turismo/alojamientos-turisticos/alojamientos_turisticos.csv", "CSV", "Descarga directa", None,
                        "Personas identificables", "Ciudad", "Registro vigente", "No se sabe", IA,
                        "Trae teléfono, correo y web de cada alojamiento (pueden ser de personas): no se copian al lago."),
    "eoh_caba": ("INDEC y Subsecretaría de Turismo de la Nación (EOH)", "Entidad pública",
                 SERIES_API + "eoh_plazas_15,eoh_viaj_loc7,eoh_viaj_loc8,eoh_viaj_loc9,eoh_pern_loc_7", "JSON", "API sin llave",
                 "Creative Commons Attribution 4.0", "Solo conteos por zona", "Ciudad (localidad CABA)", "2018-01 a 2025-08",
                 "Mensual", IA_API, "Encuesta a hoteles y parahoteles: no incluye alquileres temporarios. Rezago de ~1 año."),
    "wikimedia_pageviews": ("Wikimedia Foundation", "Plataforma colaborativa", fuente_lago["wikimedia_pageviews"]["url"], "JSON",
                            "API sin llave", "CC0", "No", "Ciudad (artículo «Buenos Aires»)", "2025-01 a 2026-08", "Mensual",
                            "Ejemplo del taller (script de la lámina), adaptado a Buenos Aires", "Agente «user»: cuenta personas, no bots. Mide interés, no opinión."),
    "ba_espacios_culturales": ("GCBA · Ministerio de Cultura", "Entidad pública",
                               CDN + "ministerio-de-cultura/espacios-culturales/espacios-culturales.csv", "CSV", "Descarga directa", None,
                               "Personas identificables", "Ciudad, por barrio", "Registro vigente", "No se sabe", IA,
                               "Trae teléfono, correo y redes de cada espacio (pueden ser de personas): no se copian al lago."),
    "ba_obras_registradas": ("GCBA · Subsecretaría de Planeamiento", "Entidad pública",
                             CDN + "subsecretaria-de-planeamiento/obras-registradas/obrasregistradas-acumulado.csv", "CSV",
                             "Descarga directa", None, "No", "Ciudad, por barrio", "2020 a 2026", "Mensual", IA,
                             "Son permisos de obra (expediente, dirección, parcela), no volúmenes construidos."),
    "ba_tejido_urbano": ("GCBA · Secretaría de Desarrollo Urbano (Secretaría de Planeamiento)", "Entidad pública",
                         CDN + "secretaria-de-desarrollo-urbano/tejido-urbano/tejido.csv", "CSV", "Descarga directa", None, "No",
                         "Ciudad, por parcela", "Hasta 2021", "Única vez", IA,
                         "Alturas por fotogrametría. Se usa el CSV (122 MB, 8 MB comprimido); el GeoJSON de ~1 GB no entra al lago (ver contrato: GeoParquet/PMTiles). 2 alturas imposibles (830 y 280 m) se descartan."),
    "ba_parcelas": ("GCBA · Secretaría de Desarrollo Urbano", "Entidad pública",
                    CDN + "secretaria-de-desarrollo-urbano/parcelas/parcelas_catastrales.zip", "Shapefile", "Descarga directa", None,
                    "No", "Ciudad, por parcela", "Catastro vigente", "No se sabe", IA,
                    "Solo se usa la tabla smp → barrio del .dbf (crudo reducido a 1 MB); la geometría no se guarda."),
    "ba_espacios_verdes": ("GCBA · Secretaría de Desarrollo Urbano", "Entidad pública",
                           CDN + "secretaria-de-desarrollo-urbano/espacios-verdes/espacio_verde_publico.csv", "CSV", "Descarga directa",
                           None, "No", "Ciudad, por barrio", "Registro vigente", "No se sabe", IA,
                           "El recurso «Espacio Verde Privado (CSV)» del mismo conjunto vino vacío (146 bytes, solo cabecera)."),
    "ba_calidad_aire": ("GCBA · Agencia de Protección Ambiental", "Entidad pública",
                        CDN + "agencia-de-proteccion-ambiental/calidad-aire/calidad_aire_2025.csv ; " + CDN + "agencia-de-proteccion-ambiental/calidad-aire/estaciones-ambientales.csv",
                        "CSV", "Descarga directa", None, "No", "4 estaciones de la Ciudad", "2025", "Diaria", IA,
                        "CO en ppm y NO2 en ppb (no µg/m³); coma decimal. Palermo sin lecturas válidas en 2025."),
    "open_meteo": ("Open-Meteo (reanálisis ERA5 del ECMWF)", "Empresa privada",
                   "https://archive-api.open-meteo.com/v1/archive?latitude=-34.6037&longitude=-58.3816&start_date=2015-01-01&end_date=2025-12-31&daily=temperature_2m_mean,temperature_2m_max,temperature_2m_min,precipitation_sum&timezone=America/Argentina/Buenos_Aires",
                   "JSON", "API sin llave", "CC BY 4.0", "No", "Un punto del centro de CABA (grilla ~10-25 km)", "2015-01 a 2025-12",
                   "Diaria", "IA (Claude Code): propuesta como fuente de clima sin llave; probada con la consulta del script",
                   "Dato de modelo, no de estación del SMN: suaviza extremos y no distingue barrios."),
    "ba_puntos_verdes": ("GCBA · Agencia de Protección Ambiental", "Entidad pública",
                         CDN + "agencia-de-proteccion-ambiental/puntos-verdes/puntos_verdes.csv ; " + CDN + "agencia-de-proteccion-ambiental/puntos-verdes/RAEEs_RAEEs.csv",
                         "CSV", "Descarga directa", None, "No", "Ciudad, por barrio", "Puntos vigentes; RAEE 2019-2026", "Mensual", IA,
                         "El primer intento rebotó por el límite de peticiones del portal. El listado solo trae 17 puntos «Con Atención». RAEE mezcla coma decimal y de miles; 7 valores ilegibles quedan sin dato."),
    "ba_establecimientos_educativos": ("GCBA · Ministerio de Educación", "Entidad pública",
                                       CDN + "ministerio-de-educacion/establecimientos-educativos/establecimientos_educativos.csv", "CSV",
                                       "Descarga directa", None, "No", "Ciudad, por barrio", "Registro vigente", "No se sabe", IA,
                                       "Una fila es una oferta (nivel/modalidad), no un edificio. Sin matrícula."),
    "ba_hospitales": ("GCBA · Ministerio de Salud", "Entidad pública", CDN + "ministerio-de-salud/hospitales/hospitales.csv", "CSV",
                      "Descarga directa", None, "No", "Ciudad, por barrio", "Registro vigente", "No se sabe", IA,
                      "Teléfono y fax institucionales: no se copian. Sin camas ni atenciones."),
    "ba_cesac": ("GCBA · Ministerio de Salud", "Entidad pública",
                 CDN + "ministerio-de-salud/centros-salud-accion-comunitaria-cesac/centros_salud_nivel_1_cesac.csv", "CSV",
                 "Descarga directa", None, "No", "Ciudad, por barrio", "Registro vigente", "No se sabe", IA,
                 "Teléfono y web institucionales: no se copian."),
    "ba_subte_estaciones": ("Subterráneos de Buenos Aires (SBASE)", "Entidad pública",
                            CDN + "sbase/subte-estaciones/estaciones_de_subte.csv ; " + CDN + "sbase/subte-estaciones/viajes_anual.csv ; " + CDN + "sbase/subte-estaciones/red_subte.geojson",
                            "CSV", "Descarga directa", None, "No", "Red de subte (6 líneas)", "2013 a 2020 (pasajeros anuales)", "Anual", IA,
                            "viajes_anual.csv tiene 48 filas año×línea: un error previo las contaba como 48 estaciones (son 90 pares estación-línea)."),
    "ba_subte_molinetes": ("Subterráneos de Buenos Aires (SBASE)", "Entidad pública",
                           CDN + "sbase/subte-viajes-molinetes/molinetes-2025.zip", "CSV", "Descarga directa", None, "No",
                           "Red de subte y Premetro, por molinete", "2025-01 a 2025-12", "Mensual", IA,
                           "ZIP de 82 MB con 26 CSV cada 15 min. Agosto mezcla fechas d/m y m/d; marzo trae «;» sobrantes; enero, líneas vacías; codificación variable; un molinete «Prueba» excluido."),
    "ba_ecobici_estaciones": ("GCBA · Secretaría de Transporte y Obras Públicas", "Entidad pública",
                              CDN + "transporte-y-obras-publicas/estaciones-bicicletas-publicas/nuevas-estaciones-bicicletas-publicas.csv", "CSV",
                              "Descarga directa", None, "No", "Ciudad", "Registro vigente", "No se sabe", IA,
                              "El campo barrio viene vacío en las 407 estaciones: se asigna por punto en polígono."),
    "ba_colectivos_paradas": ("GCBA · Secretaría de Transporte y Obras Públicas", "Entidad pública",
                              CDN + "transporte-y-obras-publicas/colectivos-paradas/paradas-de-colectivo.csv", "CSV", "Descarga directa",
                              None, "No", "Ciudad", "Registro vigente", "No se sabe", IA,
                              "Barrios con erratas («VILLA ORTUZA», «VELEZ SARFIELD»): se asigna por punto en polígono."),
    "ba_sensores_transito": ("GCBA · Secretaría de Transporte y Obras Públicas", "Entidad pública",
                             CDN + "transporte-y-obras-publicas/sensores-transito/sensores.csv", "CSV", "Descarga directa", None, "No",
                             "Ciudad, por comuna", "Inventario vigente", "No se sabe", IA, "Solo inventario de sensores: las lecturas están en la API de Transporte, hoy suspendida."),
    "ba_tmi": ("GCBA · Agencia de Protección Ambiental", "Entidad pública",
               CDN + "agencia-de-proteccion-ambiental/informacion-meteorologica/torres-de-monitoreo-inteligente-tmi.csv", "CSV",
               "Descarga directa", None, "No", "Ciudad, por comuna", "Inventario vigente", "No se sabe", IA,
               "Solo inventario de torres; las lecturas abiertas llegan hasta 2012."),
    "ba_data_ckan": ("GCBA · BA Data (portal de datos abiertos)", "Entidad pública", fuente_lago["ba_data_ckan"]["url"], "JSON",
                     "API sin llave", "por conjunto (cada conjunto declara la suya)", "No", "Ciudad", "Catálogo vigente", "Diaria",
                     "IA (Claude Code): identificó la API CKAN bajo /km/ y la probó con package_search",
                     "Es el inventario del portal (455 conjuntos), no un dato de la ciudad: sirve para el rastreo."),
}

# Fuentes del rastreo que NO entraron: (nombre, entidad, tipo, url, formato, cómo, licencia, personas, cobertura,
#                                     periodo, frecuencia, probado, estado, cómo la encontraron, observaciones)
NO_ENTRARON = [
    ("Sistema Único de Atención Ciudadana / BA Colaborativa (reclamos)", "GCBA · Atención Ciudadana", "Entidad pública",
     CDN + "atencion-ciudadana/sistema-unico-atencion-ciudadana/sistema-unico-de-atencion-ciudadana-2025.csv", "CSV", "Descarga directa",
     "CC-BY-2.5-AR", None, "Ciudad", "2011 a 2026", "Anual", "2026-10-06", "Descartada", SONDEO,
     "Un CSV por año de 100 a 277 MB (~2 GB en total): se descartó por peso. No se abrió el contenido, así que no se sabe si trae datos de personas."),
    ("SUACI · Solicitudes por barrio (agregado)", "GCBA · Atención Ciudadana", "Entidad pública",
     CDN + "sistema-unico-de-atencion-ciudadana/gcba_suaci_barrios.csv", "CSV", "Descarga directa", "CC-BY-2.5-AR",
     None, "Ciudad, por barrio", "No se sabe", "No se sabe", "2026-10-06", "Candidata", SONDEO,
     "Encontrada al revisar el SUACI: 2,7 MB, ya agregada por barrio. Resolvería el hueco de reclamos sin la descarga pesada. Falta abrirla y ver su periodo."),
    ("API Transporte Público (lecturas en tiempo real: subte, colectivos, EcoBici, tránsito)", "GCBA · Secretaría de Transporte",
     "Entidad pública", "https://api-transporte.buenosaires.gob.ar/", "JSON", "API con llave gratuita", "CC-BY-2.5-AR", "No",
     "Ciudad", "Tiempo real", "Tiempo real", "2026-10-06", "No responde", SONDEO,
     "El portal avisa: «todos los datasets con Formato API y GTFS están suspendidos… en revisión y corrección». Antes se había anotado que faltaban credenciales: la causa real es la suspensión."),
    ("Colectivos: GTFS", "GCBA · Secretaría de Transporte", "Entidad pública",
     "https://data.buenosaires.gob.ar/km/dataset/colectivos-gtfs", "Otro", "Descarga directa", "CC-BY-2.5-AR", "No", "Ciudad",
     "No se sabe", "No se sabe", "2026-10-06", "No responde", SONDEO,
     "Alcanzada por el mismo aviso de suspensión de los conjuntos GTFS. No se intentó la descarga."),
    ("Información Censal por Radio", "GCBA · Dirección General de Estadísticas y Censos", "Entidad pública",
     CDN + "direccion-general-de-estadisticas-y-censos/informacion-censal-por-radio/informacion-censal-por-radio-2010.csv", "CSV",
     "Descarga directa", "CC-BY-2.5-AR", "Solo conteos por zona", "Ciudad, por radio censal", "Censos 2001 y 2010", "Única vez",
     "2026-10-06", "Descartada", "IA (Claude Code): la propuso como vía a la población 2022 por barrio",
     "Solo trae radios de 2001 y 2010: no da población 2022 por barrio, que era para lo que se buscaba."),
    ("Padrón de Establecimientos Educativos", "GCBA · Ministerio de Educación", "Entidad pública",
     CDN + "ministerio-de-educacion/establecimientos-educativos/padron-establecimientos.csv", "CSV", "Descarga directa",
     "CC-BY-2.5-AR", "Personas identificables", "Ciudad", "Registro vigente", "No se sabe", "2026-10-06", "Descartada", SONDEO,
     "Trae correos (algunos personales, de Gmail/Hotmail). Se usó «Establecimientos Educativos», que no los trae."),
    ("Edificios Públicos en 3D (escaneo LIDAR)", "GCBA · Secretaría de Desarrollo Urbano", "Entidad pública",
     CDN + "secretaria-de-desarrollo-urbano/edificios-publicos-3d/cabildo.zip", "Otro", "Descarga directa", "CC-BY-2.5-AR", "No",
     "9 edificios emblemáticos", "No se sabe", "Única vez", "2026-10-06", "Descartada", SONDEO,
     "Solo 9 edificios (Cabildo, Colón, Congreso…), de 51 a 599 MB cada uno: sirve de muestra, no para la ciudad."),
    ("Área edificable: volumen", "GCBA · Secretaría de Desarrollo Urbano", "Entidad pública",
     CDN + "secretaria-de-desarrollo-urbano/area-edificable-volumen/volumen_edif.csv", "CSV", "Descarga directa", "CC-BY-2.5-AR", "No",
     "Ciudad, por parcela", "No se sabe", "No se sabe", "2026-10-06", "Descartada", SONDEO,
     "Es el volumen que permite la norma, no lo construido (38 MB). Para alturas reales se usó Tejido Urbano."),
    ("IPC Nivel General Gran Buenos Aires (INDEC)", "INDEC", "Entidad pública", SERIES_API + "145.3_INGGBAGBA_DICI_M_10", "JSON",
     "API sin llave", "Creative Commons Attribution 4.0", "No", "Gran Buenos Aires", "2016-12 a 2026-08", "Mensual", "2026-10-06",
     "Descartada", SONDEO_API, "Aparece en la búsqueda de la API con datos hasta 2026-08 (no se bajó la serie). Mide el aglomerado GBA: se prefirió el IPCBA, que mide la Ciudad."),
    ("Tasa de desempleo en hogares urbanos, CABA", "Datos Argentina (indicadores PMyE)", "Entidad pública",
     SERIES_API + "obras_13_08_02", "JSON", "API sin llave", None, "Solo conteos por zona", "Ciudad", "2019-10 a 2023-04",
     "Trimestral", "2026-10-06", "Descartada", SONDEO_API, "La búsqueda de la API la muestra cortada en 2023-04: vieja para el sistema. Se usó el empleo registrado (SIPA). Licencia no revisada."),
]

# ---------------------------------------------------------------- 2. Diccionario
PREG = {
    "gente": "P1 · ¿Cuánta gente vive en cada barrio y qué servicios de salud, educación y cultura tiene cerca?",
    "plata": "P2 · ¿Cómo gasta la Ciudad su presupuesto y cuánto vale en pesos de hoy?",
    "seguridad": "P3 · ¿Qué pasa con la seguridad en cada barrio?",
    "economia": "P4 · ¿Cuánto pesan el turismo, el empleo y los precios?",
    "suelo": "P5 · ¿Qué hay construido, cuánto verde queda y cómo está el ambiente?",
    "movilidad": "P6 · ¿Cómo se mueve la gente por la Ciudad?",
    "metodo": "Fuentes y método · ¿de dónde sale cada cifra y qué tan confiable es?",
    "cruce": "Cruza P1, P3, P5 y P6 · ¿qué barrios se parecen y en qué?",
}
TEMA_PREG = {"demografia": "gente", "salud": "gente", "educacion": "gente", "cultura": "gente", "atlas_distritos": "gente", "atlas_zonas": "gente",
             "municipio": "plata", "seguridad": "seguridad", "economia": "economia", "turismo": "economia", "escucha": "economia",
             "edificios": "suelo", "espacios_verdes": "suelo", "calidad_aire": "suelo", "clima": "suelo", "residuos": "suelo",
             "sensores": "movilidad", "geoinfraestructura": "movilidad", "subte_molinetes": "movilidad", "movilidad": "movilidad",
             "catalogo": "metodo", "correlaciones": "cruce"}

# (tema, clave normalizada): (nombre original en la fuente, qué significa, transformación)
D = {
    ("atlas_distritos", "registros"): ("features del GeoJSON", "Cantidad de barrios de la Ciudad", "Conteo de polígonos"),
    ("atlas_distritos", "area_km2_por_barrio"): ("area (m²)", "Superficie de cada barrio", "m² → km², 2 decimales"),
    ("atlas_distritos", "comuna_por_barrio"): ("comuna", "Comuna a la que pertenece cada barrio", "Se copia tal cual"),
    ("atlas_zonas", "registros"): ("features del GeoJSON", "Cantidad de comunas", "Conteo de polígonos"),
    ("atlas_zonas", "area_km2_por_comuna"): ("area (m²)", "Superficie de cada comuna", "m² → km², 1 decimal"),
    ("calidad_aire", "estaciones_ambientales"): ("nombre (estaciones.csv)", "Estaciones de monitoreo del aire de la Ciudad", "Conteo de filas"),
    ("calidad_aire", "lecturas_horarias"): ("fecha, hora", "Horas con registro en el año", "Conteo de filas fecha-hora"),
    ("calidad_aire", "promedio_co_por_estacion"): ("co_<estación>", "Monóxido de carbono promedio del año en cada estación", "Coma decimal → punto; promedio sin vacíos; ppm (no µg/m³)"),
    ("calidad_aire", "promedio_no2_por_estacion"): ("no2_<estación>", "Dióxido de nitrógeno promedio del año en cada estación", "Coma decimal → punto; promedio sin vacíos; ppb"),
    ("calidad_aire", "promedio_pm10_por_estacion"): ("pm10_<estación>", "Material particulado PM10 promedio del año en cada estación", "Coma decimal → punto; promedio sin vacíos"),
    ("catalogo", "conjuntos"): ("result.count", "Conjuntos de datos publicados en BA Data", "Se copia el total de package_search"),
    ("clima", "dias_32c_<año>"): ("temperature_2m_max", "Días del año con máxima de 32 °C o más", "Conteo de días diarios ≥ 32"),
    ("clima", "maxima_<año>"): ("temperature_2m_max", "Máxima diaria más alta del año", "Máximo de la serie diaria"),
    ("clima", "precipitacion_<año>"): ("precipitation_sum", "Lluvia total del año", "Suma de la serie diaria"),
    ("clima", "temperatura_media_<año>"): ("temperature_2m_mean", "Temperatura media del año", "Promedio de la serie diaria"),
    ("clima", "dias_32c_anual"): ("temperature_2m_max", "Días con máxima ≥ 32 °C, año por año", "Conteo por año"),
    ("clima", "precipitacion_anual"): ("precipitation_sum", "Lluvia total de cada año", "Suma por año"),
    ("clima", "precipitacion_mensual"): ("precipitation_sum", "Lluvia total de cada mes", "Suma por mes, 1 decimal"),
    ("clima", "temperatura_media_anual"): ("temperature_2m_mean", "Temperatura media de cada año", "Promedio por año"),
    ("clima", "temperatura_media_mensual"): ("temperature_2m_mean", "Temperatura media de cada mes", "Promedio por mes, 1 decimal"),
    ("correlaciones", "barrios_comparados"): ("—", "Barrios con población y área para comparar", "Intersección de barrios entre temas"),
    ("correlaciones", "pares_fuertes"): ("—", "Pares de indicadores con asociación fuerte (|r| ≥ 0,5)", "Conteo sobre la matriz de Spearman"),
    ("correlaciones", "spearman"): ("—", "Correlación de rangos entre indicadores por barrio (asociación, no causa)", "Tasas por habitante y densidad; Spearman con rangos empatados"),
    ("correlaciones", "spearman_fuertes"): ("—", "Subconjunto de pares con |r| ≥ 0,5", "Filtro de la matriz"),
    ("cultura", "registros"): ("fid", "Espacios culturales registrados", "Conteo de filas (sin contactos)"),
    ("cultura", "por_barrio"): ("BARRIO", "Espacios culturales en cada barrio", "Mayúsculas; conteo"),
    ("cultura", "por_funcion_principal"): ("FUNCION_PRINCIPAL", "Espacios culturales por tipo (biblioteca, teatro…)", "Conteo"),
    ("demografia", "poblacion_barrios"): ("POBLACION", "Habitantes de la Ciudad sumando los barrios (Censo 2010)", "Suma de barrios"),
    ("demografia", "poblacion_comunas"): ("Anexo 1, cuadro C.1 (PDF)", "Habitantes de la Ciudad sumando las comunas (Censo 2022)", "Transcrito del PDF; la suma se compara con el total publicado"),
    ("demografia", "poblacion_por_barrio"): ("BARRIO, POBLACION", "Habitantes de cada barrio en el Censo 2010", "Se copia; nombres en mayúsculas"),
    ("demografia", "poblacion_por_comuna"): ("Anexo 1, cuadro C.1 (PDF)", "Habitantes de cada comuna en el Censo 2022", "Transcrito del PDF a ingesta/datos/censo2022_comunas.csv; la suma = total oficial"),
    ("economia", "asalariados_privados_ultimo_mes"): ("serie 154.2_CABA_S_0_0_4", "Empleos privados registrados en CABA el último mes", "Miles × 1000"),
    ("economia", "inflacion_<año>"): ("serie 193.2_NIVEL_GENERAL_2021_0_13_2", "Inflación de la Ciudad diciembre a diciembre", "(IPC dic / IPC dic año anterior − 1) × 100"),
    ("economia", "inflacion_interanual"): ("serie 193.2_NIVEL_GENERAL_2021_0_13_2", "Inflación de los últimos 12 meses", "(IPC último / IPC hace 12 meses − 1) × 100"),
    ("economia", "ipcba_promedio_<año>"): ("serie 193.2_NIVEL_GENERAL_2021_0_13_2", "Índice de precios promedio del año (para deflactar)", "Promedio de 12 meses"),
    ("economia", "ipcba_ultimo_mes"): ("serie 193.2_NIVEL_GENERAL_2021_0_13_2", "Índice de precios de la Ciudad del último mes", "Se copia"),
    ("economia", "asalariados_privados"): ("serie 154.2_CABA_S_0_0_4", "Empleo privado registrado mes a mes", "Fecha AAAA-MM-DD → AAAA-MM"),
    ("economia", "ipcba"): ("serie 193.2_NIVEL_GENERAL_2021_0_13_2", "Índice de precios de la Ciudad mes a mes", "Fecha AAAA-MM-DD → AAAA-MM"),
    ("edificios", "altura_maxima"): ("altura", "Altura del volumen edificado más alto", "Máximo, sin alturas > 250 m"),
    ("edificios", "altura_mediana_parcela"): ("altura", "Altura típica de una parcela de la Ciudad", "Máximo por parcela, luego mediana"),
    ("edificios", "parcelas_cruzadas_con_barrio"): ("smp, barrio", "Parcelas edificadas a las que se les encontró barrio", "Cruce por smp normalizado; si falta, por manzana"),
    ("edificios", "parcelas_edificadas"): ("smp", "Parcelas con al menos un volumen edificado", "smp sin ceros a la izquierda; Ñ reparada"),
    ("edificios", "registros"): ("id_registr", "Permisos de obra registrados", "Conteo de filas"),
    ("edificios", "volumenes_descartados"): ("altura", "Volúmenes con altura imposible (0 o > 250 m)", "Conteo de excluidos"),
    ("edificios", "volumenes_edificados"): ("id", "Volúmenes edificados en el tejido", "Conteo de filas"),
    ("edificios", "altura_mediana_por_barrio"): ("altura, barrio", "Altura típica de las parcelas de cada barrio", "Mediana de la altura máxima por parcela"),
    ("edificios", "parcelas_edificadas_por_barrio"): ("smp, barrio", "Parcelas edificadas en cada barrio", "Conteo tras el cruce"),
    ("edificios", "parcelas_por_rango_altura"): ("altura", "Parcelas según su altura (≈ pisos)", "Rangos 0-9, 9-15, 15-30, 30-60, >60 m"),
    ("edificios", "parcelas_por_via_de_cruce"): ("smp", "Cómo se le encontró barrio a cada parcela", "parcela / manzana / sin cruce"),
    ("edificios", "pct_parcelas_mas_30m_por_barrio"): ("altura, barrio", "Porcentaje de parcelas con más de 30 m (≈ 10 pisos) por barrio", "Proporción × 100"),
    ("edificios", "por_anio"): ("fecha", "Permisos de obra por año", "Año extraído de dd/mm/aaaa"),
    ("edificios", "por_barrio"): ("barrio", "Permisos de obra en cada barrio", "Mayúsculas; conteo"),
    ("edificios", "volumenes_por_tipo"): ("tipo", "Volúmenes edificados por tipo", "Conteo"),
    ("educacion", "estatales"): ("ges", "Establecimientos de gestión estatal", "Conteo de ges = Estatal"),
    ("educacion", "registros"): ("id", "Ofertas educativas (nivel/modalidad por establecimiento)", "Conteo de filas"),
    ("educacion", "por_barrio"): ("bar", "Ofertas educativas en cada barrio", "Mayúsculas sin tildes; conteo"),
    ("educacion", "por_gestion"): ("ges", "Ofertas por tipo de gestión (estatal/privada)", "Conteo"),
    ("educacion", "por_nivel"): ("nen_mde", "Ofertas por nivel educativo", "Conteo"),
    ("escucha", "vistas_ultimo_mes"): ("views", "Vistas del artículo «Buenos Aires» hechas por personas en el último mes", "Último mes de la serie; AAAAMMDD → AAAA-MM"),
    ("escucha", "vistas_mensuales"): ("views", "Vistas mensuales del artículo «Buenos Aires»", "AAAAMMDD → AAAA-MM"),
    ("espacios_verdes", "registros"): ("id", "Espacios verdes públicos", "Conteo de filas"),
    ("espacios_verdes", "superficie_m2"): ("area", "Superficie total de verde público", "Suma, redondeo a m²"),
    ("espacios_verdes", "m2_por_barrio"): ("area, barrio", "Metros cuadrados de verde público en cada barrio", "Suma por barrio"),
    ("espacios_verdes", "m2_por_clasificacion"): ("area, clasificac", "Metros cuadrados por tipo (parque, plaza, plazoleta…)", "Suma por tipo"),
    ("espacios_verdes", "por_barrio"): ("barrio", "Espacios verdes públicos en cada barrio", "Mayúsculas sin tildes; conteo"),
    ("geoinfraestructura", "estaciones_distintas"): ("estacion", "Estaciones de subte con nombre distinto", "Conteo de nombres únicos"),
    ("geoinfraestructura", "estaciones_linea"): ("id", "Pares estación-línea (las de combinación cuentan una vez por línea)", "Conteo de filas"),
    ("geoinfraestructura", "lineas"): ("linea", "Líneas de subte", "Conteo de valores únicos"),
    ("geoinfraestructura", "pasajeros_linea_<X>"): ("year, linea, total", "Pasajeros anuales de cada línea (A, B, C, D, E, H)", "Una serie por línea"),
    ("geoinfraestructura", "estaciones_por_linea"): ("linea", "Estaciones de cada línea", "Conteo"),
    ("movilidad", "estaciones_ecobici"): ("id", "Estaciones de bicicletas públicas", "Conteo de filas"),
    ("movilidad", "lineas_colectivo"): ("L1 … L6", "Líneas de colectivo distintas con parada en la Ciudad", "Unión de L1-L6"),
    ("movilidad", "paradas_colectivo"): ("—", "Paradas de colectivo", "Conteo de filas"),
    ("movilidad", "ecobici_por_barrio"): ("latitud, longitud", "Estaciones EcoBici en cada barrio", "Barrio por punto en polígono (el campo barrio viene vacío)"),
    ("movilidad", "paradas_por_barrio"): ("coord_X, coord_Y", "Paradas de colectivo en cada barrio", "Barrio por punto en polígono (el campo trae erratas)"),
    ("municipio", "devengado_<año>"): ("Devengado / Devengado_Trim2_CONT", "Gasto ya comprometido y realizado (devengado)", "Punto de miles → entero; suma"),
    ("municipio", "devengado_<año>_a_precios_2026-08"): ("Devengado + IPCBA", "Gasto 2025 expresado en pesos de agosto de 2026", "× IPCBA último / IPCBA promedio 2025"),
    ("municipio", "ejecucion_<año>"): ("Devengado, Vigente", "Qué parte del presupuesto vigente ya se gastó", "devengado / vigente × 100"),
    ("municipio", "registros"): ("—", "Filas del archivo de ejecución presupuestaria", "Conteo sin filas vacías"),
    ("municipio", "sancionado_<año>"): ("Sanción", "Presupuesto aprobado por la Legislatura", "Punto de miles → entero; suma"),
    ("municipio", "vigente_<año>"): ("Vigente / Vigente_Trim2_CONT", "Presupuesto con las modificaciones del año", "Punto de miles → entero; suma"),
    ("municipio", "devengado_por_finalidad_<año>"): ("Desc_Fin", "Gasto devengado por finalidad (social, seguridad…)", "Suma por finalidad"),
    ("municipio", "devengado_por_jurisdiccion_<año>"): ("Desc_Jur", "Gasto devengado por ministerio u organismo", "Suma por jurisdicción"),
    ("municipio", "devengado_por_ubicacion_<año>"): ("Desc_Geo", "Gasto devengado por ubicación geográfica imputada", "Suma; no es gasto «en» el barrio"),
    ("residuos", "puntos_verdes"): ("id", "Puntos verdes de reciclaje listados", "Conteo de filas"),
    ("residuos", "raee_kg_<año>"): ("kilos", "Kilos de aparatos eléctricos y electrónicos recibidos en el año", "Coma decimal o de miles por patrón; ilegibles fuera"),
    ("residuos", "raee_kg_anual"): ("año, kilos", "Kilos de RAEE recibidos cada año", "Suma por año"),
    ("residuos", "por_barrio"): ("barrio", "Puntos verdes en cada barrio", "Mayúsculas sin tildes; conteo"),
    ("residuos", "por_tipo"): ("tipo", "Puntos verdes por tipo", "Conteo"),
    ("salud", "cesac"): ("id", "Centros de Salud y Acción Comunitaria (atención primaria)", "Conteo de filas"),
    ("salud", "registros"): ("fna", "Hospitales públicos de la Ciudad", "Conteo de filas"),
    ("salud", "cesac_por_area_programatica"): ("area_progr", "CESAC por hospital del que dependen", "Conteo"),
    ("salud", "cesac_por_barrio"): ("barrio", "CESAC en cada barrio", "Mayúsculas sin tildes; conteo"),
    ("salud", "por_barrio"): ("bar", "Hospitales en cada barrio", "Mayúsculas sin tildes; conteo"),
    ("salud", "por_especialidad"): ("esp", "Hospitales por especialidad", "Conteo"),
    ("salud", "por_tipo"): ("gna", "Hospitales por tipo (agudos, especializado, niños)", "Conteo"),
    ("seguridad", "delitos_<año>"): ("cantidad", "Delitos registrados en el año", "Suma de cantidad"),
    ("seguridad", "meses_con_datos_<año>"): ("mes", "Meses con datos en el año (12 = año completo)", "Conteo de meses distintos"),
    ("seguridad", "delitos_mensuales_<año>"): ("mes, cantidad", "Delitos registrados mes a mes", "Mes en letras → AAAA-MM; suma"),
    ("seguridad", "por_barrio_<año>"): ("barrio, cantidad", "Delitos registrados en cada barrio", "Mayúsculas; suma"),
    ("seguridad", "por_tipo_<año>"): ("tipo, cantidad", "Delitos por tipo (robo, hurto, lesiones…)", "Suma"),
    ("sensores", "ba_sensores_transito"): ("id_sensor", "Sensores de tránsito instalados", "Conteo de filas"),
    ("sensores", "ba_tmi"): ("tmi", "Torres de monitoreo inteligente", "Conteo de filas"),
    ("sensores", "ba_sensores_transito_por_comuna"): ("comuna", "Sensores de tránsito por comuna", "Conteo"),
    ("sensores", "ba_tmi_por_comuna"): ("comuna", "Torres de monitoreo por comuna", "Conteo"),
    ("subte_molinetes", "lineas"): ("LINEA", "Líneas con molinetes (A-H y Premetro)", "«LineaA» → «A»; sin «Prueba»"),
    ("subte_molinetes", "meses_con_datos_<año>"): ("nombre del archivo AAAAMM", "Meses con datos (12 = año completo)", "Mes tomado del nombre del archivo"),
    ("subte_molinetes", "pasajeros_subte_<año>"): ("pax_TOTAL", "Validaciones en molinetes del subte en el año (sin Premetro)", "Suma de pax_TOTAL"),
    ("subte_molinetes", "pasajeros_linea_<X>"): ("pax_TOTAL, LINEA", "Validaciones mensuales de cada línea (A, B, C, D, E, H, PM)", "Suma por línea y mes"),
    ("subte_molinetes", "pasajeros_por_barrio_<año>"): ("pax_TOTAL + capa geo", "Validaciones en las estaciones de cada barrio", "Estación → capa geo → barrio por punto en polígono"),
    ("subte_molinetes", "pasajeros_por_estacion_<año>"): ("ESTACION", "Validaciones por estación, con el nombre del crudo", "Suma; #N/D y NULL → sin dato"),
    ("subte_molinetes", "pasajeros_por_estacion_geo_<año>"): ("ESTACION", "Validaciones por estación, con el nombre de la capa geo", "Nombres normalizados con alias revisados a mano"),
    ("subte_molinetes", "sin_estacion_geo_<año>"): ("ESTACION", "Validaciones que no se pudieron ubicar en una estación", "Lo que no cruzó con la capa geo"),
    ("turismo", "alojamientos"): ("id", "Alojamientos turísticos registrados", "Conteo (sin contactos)"),
    ("turismo", "ocupacion_plazas_ultimo_mes"): ("serie eoh_plazas_15", "Porcentaje de plazas hoteleras ocupadas el último mes", "Fracción × 100"),
    ("turismo", "pernoctaciones_ultimo_mes"): ("serie eoh_pern_loc_7", "Noches pasadas por viajeros el último mes", "Se copia"),
    ("turismo", "viajeros_<año>"): ("serie eoh_viaj_loc7", "Viajeros hospedados en el último año completo", "Suma de 12 meses"),
    ("turismo", "viajeros_no_residentes_ultimo_mes"): ("serie eoh_viaj_loc9", "Viajeros extranjeros hospedados el último mes", "Se copia"),
    ("turismo", "viajeros_residentes_ultimo_mes"): ("serie eoh_viaj_loc8", "Viajeros argentinos hospedados el último mes", "Se copia"),
    ("turismo", "viajeros_total_ultimo_mes"): ("serie eoh_viaj_loc7", "Viajeros hospedados el último mes", "Se copia"),
    ("turismo", "ocupacion_plazas"): ("serie eoh_plazas_15", "Ocupación hotelera mes a mes", "Fracción × 100"),
    ("turismo", "pernoctaciones"): ("serie eoh_pern_loc_7", "Noches pasadas por viajeros mes a mes", "AAAA-MM-DD → AAAA-MM"),
    ("turismo", "viajeros_no_residentes"): ("serie eoh_viaj_loc9", "Viajeros extranjeros mes a mes", "AAAA-MM-DD → AAAA-MM"),
    ("turismo", "viajeros_residentes"): ("serie eoh_viaj_loc8", "Viajeros argentinos mes a mes", "AAAA-MM-DD → AAAA-MM"),
    ("turismo", "viajeros_total"): ("serie eoh_viaj_loc7", "Viajeros hospedados mes a mes", "AAAA-MM-DD → AAAA-MM"),
    ("turismo", "por_tipo"): ("tipo", "Alojamientos por tipo de registro", "Conteo"),
}


def norm(k):
    return re.sub(r"_linea_[A-Z][A-Za-z]*$", "_linea_<X>", re.sub(r"(?<=_)\d{4}(?=_|$)", "<año>", k))


def tipo_y_rango(valores, unidad):
    u = unidad.lower()
    if "spearman" in u:
        return "Decimal", "−1 a 1"
    if "%" in u:
        return "Decimal", "0 a 100"
    if "°c" in u:
        return "Decimal", "−10 a 45"
    entero = all(isinstance(v, int) for v in valores)
    return ("Entero" if entero else "Decimal"), "≥ 0"


filas_dic = []
# Estructura del contrato: común a todos los archivos del lago
COMUNES = [
    ("tema", "—", "Nombre del tema; coincide con el nombre del archivo", "Texto", "—", "seguridad", "uno por archivo", "Lo escribe el script"),
    ("probado", "—", "Día en que la fuente respondió (fecha de la carpeta del crudo)", "Fecha", "AAAA-MM-DD", "2026-10-06", "una fecha", "Día de la descarga, sin hora"),
    ("fuentes[] (id, nombre, url, estado, licencia)", "metadatos del portal", "Ficha corta de cada fuente usada en el archivo", "Texto", "—", "ba_delitos · CC-BY-2.5-AR", "estado: vivo o declarado", "id propio; licencia del portal"),
    ("cifras.<clave>.valor / unidad / vigencia / fuente", "—", "Cada cifra con su valor, unidad, vigencia (AAAA o AAAA-MM) y fuente", "Decimal", "la de cada cifra", "158838 · delitos registrados · 2024", "valor ≥ 0 salvo coeficientes", "El verificador exige los 4 campos"),
    ("series.<clave>.puntos", "—", "Serie temporal: pares [periodo, número]", "Decimal", "la de cada serie", "[\"2024-01\", 13626]", "periodo AAAA o AAAA-MM", "Fechas normalizadas a AAAA-MM"),
    ("desgloses.<clave>.filas", "—", "Reparto de una cifra: pares [categoría, número], ordenados de mayor a menor", "Decimal", "la de cada desglose", "[\"PALERMO\", 12794]", "categorías de la fuente", "Agregado: nunca un registro individual"),
    ("nota", "—", "Advertencias para leer bien las cifras (cobertura, rezagos, sesgos)", "Texto", "—", "Son hechos registrados, no tasas…", "texto libre escrito por el grupo", "—"),
]
for campo, orig, sig, tipo, unidad, ej, rango, transf in COMUNES:
    filas_dic.append(["todas", "lago/*.json (todos)", campo, orig, sig, tipo, unidad, ej, "No", rango, "Pública", transf, PREG["metodo"]])

fid = {}  # id lago → F##
for i, k in enumerate(INTEGRADAS, 1):
    fid[k] = f"F{i:02d}"
faltan = []
for tema, j in sorted(lago.items()):
    for sec in ("cifras", "series", "desgloses"):
        grupos = {}
        for k, v in j.get(sec, {}).items():
            grupos.setdefault(norm(k), []).append((k, v))
        for nk, items in grupos.items():
            k0, v0 = items[0]
            if (tema, nk) not in D:
                faltan.append((tema, nk))
                continue
            orig, sig, transf = D[(tema, nk)]
            fuentes = sorted({fid[v["fuente"]] for _, v in items})
            if sec == "cifras":
                vals = [v["valor"] for _, v in items]
                ej = f'{v0["valor"]} ({v0["vigencia"]})'
                nombre = nk if len(items) > 1 else k0
            else:
                pares = [p for _, v in items for p in (v.get("puntos") or v.get("filas"))]
                vals = [p[1] for p in pares]
                ej = json.dumps((v0.get("puntos") or v0.get("filas"))[0], ensure_ascii=False)
                nombre = nk if len(items) > 1 else k0
                sig = sig + (" (serie [periodo, valor])" if sec == "series" else " (desglose [categoría, valor])")
            if len(items) > 1:
                nombre += "  ⟵ " + ", ".join(k for k, _ in items)
            tipo, rango = tipo_y_rango(vals, v0["unidad"])
            filas_dic.append([", ".join(fuentes), f"lago/{tema}.json · {sec}", nombre, orig, sig, tipo, v0["unidad"], ej, "No",
                              rango, "Pública", transf, PREG[TEMA_PREG[tema]]])
    for extra, (sig, transf) in {"campos": ("Nombres de columnas del crudo, para trazabilidad", "Se copian de la cabecera"),
                                  "recurso": ("Nombre, URL y formato del recurso descargado", "Se copian del portal"),
                                  "conjuntos": ("Inventario de conjuntos del portal: id, título, licencia, actualizado, formatos", "Se copian de package_search")}.items():
        if extra in j:
            filas_dic.append([fid[j["fuentes"][0]["id"]], f"lago/{tema}.json", extra, "metadatos del portal", sig, "Texto", "—",
                              str(j[extra] if not isinstance(j[extra], list) else j[extra][0])[:80], "No", "—", "Pública", transf, PREG["metodo"]])
if faltan:
    raise SystemExit(f"Campos del lago sin descripción en el diccionario: {faltan}")

GEO = {  # capa: {propiedad: (original, significado, tipo, ejemplo/transf)}
    "barrios": {"id": ("id", "Identificador del barrio", "Entero"), "nombre": ("nombre", "Nombre del barrio", "Texto"),
                "comuna": ("comuna", "Comuna del barrio", "Entero"), "geometry": ("geometry", "Polígono del barrio (WGS84, simplificado ~10 m)", "Coordenada / geometría")},
    "comunas": {"id": ("id", "Identificador de la comuna", "Entero"), "comuna": ("comuna", "Número de comuna", "Entero"),
                "barrios": ("barrios", "Barrios que forman la comuna", "Texto"), "geometry": ("geometry", "Polígono de la comuna (WGS84, simplificado)", "Coordenada / geometría")},
    "subte_estaciones": {"estacion": ("estacion", "Nombre de la estación", "Texto"), "linea": ("linea", "Línea (A, B, C, D, E, H)", "Código"),
                         "geometry": ("geometry", "Punto de la estación (WGS84)", "Coordenada / geometría")},
    "subte_lineas": {"id": ("id", "Identificador del tramo", "Entero"), "nombre": ("nombre", "Nombre de la línea del tramo", "Texto"),
                     "geometry": ("geometry", "Trazado del tramo (WGS84, simplificado)", "Coordenada / geometría")},
    "subte_pasajeros": {"estacion": ("estacion", "Estación (como en subte_estaciones)", "Texto"), "linea": ("linea", "Línea", "Código"),
                        "barrio": ("— (calculado)", "Barrio donde está la estación", "Texto"),
                        "pasajeros_2025": ("pax_TOTAL", "Validaciones en molinetes de la estación en 2025", "Entero"),
                        "geometry": ("geometry", "Punto de la estación (WGS84)", "Coordenada / geometría")},
}
for capa, props in GEO.items():
    g = geo[capa]
    f0 = g["features"][0]
    for p, (orig, sig, tipo) in props.items():
        ej = f0["geometry"]["type"] if p == "geometry" else f0["properties"][p]
        transf = ("Douglas-Peucker ~10 m, 5 decimales" if p == "geometry" and g["gemelo"].get("simplificacion", "").startswith("Douglas") and
                  f0["geometry"]["type"] != "Point" else "Barrio por punto en polígono" if (capa, p) == ("subte_pasajeros", "barrio")
                  else "Suma 2025 por estación" if p.startswith("pasajeros") else "Se copia")
        filas_dic.append([fid[g["gemelo"]["fuente"]], f"lago/geo/{capa}.geojson", p, orig, sig, tipo,
                          "grados WGS84" if p == "geometry" else ("pasajeros" if p.startswith("pasajeros") else "—"), str(ej)[:60],
                          "No", "dentro de la Ciudad" if p == "geometry" else ("≥ 0" if tipo == "Entero" else "—"), "Pública", transf,
                          PREG["movilidad" if capa.startswith("subte") else "gente"]])

# ---------------------------------------------------------------- 3. Bitácora IA (sesión anterior: Claude Sonnet 5.5; esta: Claude Opus 5.5)
S1, S2 = "Claude Code (Claude Sonnet 5.5)", "Claude Code (Claude Opus 5.5)"
BITACORA = [
    (S1, "Encontrar el portal de datos abiertos de Buenos Aires y su API", "BA Data en CKAN; la API está bajo /km/api/3/action",
     "package_search respondió con 455 conjuntos", "Correcto", "ba_data_ckan", "Antes de confiar en una API hay que ver que el endpoint sigue vigente: la ruta cambió (/km/)."),
    (S1, "Contar las estaciones de subte", "Primero dio 48 estaciones", "Se abrió el CSV: 48 eran filas año×línea de pasajeros; las estaciones son 90 pares estación-línea (88 nombres)",
     "Incorrecto", "ba_subte_estaciones", "Un número redondo y verosímil también puede ser el conteo de otra cosa: mirar qué es cada fila."),
    (S1, "Curar la calidad del aire", "Etiquetó todo en µg/m³", "Se revisó la documentación del conjunto: CO va en ppm y NO2 en ppb",
     "Incorrecto", "ba_calidad_aire", "Las unidades se comprueban en la fuente, no se suponen por costumbre."),
    (S1, "Contar filas del presupuesto", "34.836 filas", "Se encontró una fila vacía contada como dato: 34.835", "Parcial", "ba_presupuesto",
     "Los conteos deben ignorar filas vacías; ahora lo hace leer_csv."),
    (S1, "Reemplazar la población por el Censo 2022", "Encontró el informe 1877 de Estadística Ciudad con población por comuna",
     "Se transcribió el cuadro C.1 y la suma de las 15 comunas dio el total oficial (3.121.707)", "Correcto", "dgeyc_censo2022",
     "Un PDF sirve si se marca como «declarado» y se verifica contra un total."),
    (S1, "Conseguir población 2022 por barrio", "Propuso los radios censales de BA Data como vía", "Al abrir el conjunto: solo hay radios 2001 y 2010",
     "Incorrecto", "F35", "Que exista un conjunto con el nombre correcto no quiere decir que tenga el año que se necesita."),
    (S1, "Buscar una fuente de edificios en 3D", "Tejido Urbano (alturas), Edificios Públicos 3D (LIDAR) y Área edificable",
     "Se miraron tamaños y campos con descargas parciales: solo el Tejido trae alturas de toda la ciudad", "Correcto", "ba_tejido_urbano",
     "«3D» puede ser alturas, nubes LIDAR o volumen permitido: tres cosas distintas."),
    (S1, "Explicar por qué los sensores en vivo no entraban", "Dijo que la API pedía credenciales",
     "El 2026-10-06 la ficha del portal dice que las API y GTFS están suspendidas", "Parcial", "F33",
     "La causa de un hueco también se verifica: el motivo anotado era otro."),
    (S2, "Buscar ocupación hotelera de CABA", "No estaba en BA Data; encontró la EOH (INDEC/Turismo) en la API de series de datos.gob.ar",
     "Se bajaron las series: residentes + no residentes = total en cada mes", "Correcto", "eoh_caba",
     "Salir del portal de la ciudad bajó la dependencia de un solo proveedor."),
    (S2, "Sumar los montos del presupuesto y describir el gasto por comuna", "Escribió que el gasto se concentraba en la Comuna 1",
     "Los datos mostraron que la comuna con más gasto imputado es la 4", "Incorrecto", "ba_presupuesto",
     "Una nota escrita antes de ver los datos es una suposición: se escribe después."),
    (S2, "Leer los 26 CSV de molinetes 2025", "Primer parser con fecha d/m", "Falló: agosto mezcla d/m y m/d; marzo trae «;» sobrantes; enero 82 mil líneas vacías; un molinete «Prueba»",
     "Parcial", "ba_subte_molinetes", "Validar cada fila contra el mes del archivo destapó errores que la suma no mostraba."),
    (S2, "Cruzar alturas del Tejido con barrios", "Primer resultado: altura máxima 830 m", "Imposible (la torre más alta ronda los 235 m): eran 2 registros especiales; además una Ñ doblemente codificada rompía el cruce",
     "Parcial", "ba_tejido_urbano", "Comparar el máximo con un dato conocido del mundo real es una prueba barata."),
    (S2, "Nombrar estaciones de molinetes como en la capa geo", "Normalización + alias", "Se revisaron a mano las 104 correspondencias: todas correctas; faltaba «C. Pellegrini»",
     "Correcto", "ba_subte_molinetes", "Las correspondencias automáticas se revisan una por una cuando son pocas."),
    (S2, "Proponer fuentes para llegar a 30 sin inventar", "IPCBA, SIPA, Open-Meteo, CESAC, EcoBici, paradas y puntos verdes",
     "Cada una se consultó con el script real; puntos verdes rebotó por límite de peticiones y se reintentó", "Correcto", "ipcba",
     "Contar conjuntos no es contar proveedores: 24 de 30 siguen siendo de BA Data."),
    (S2, "Asignar barrio a EcoBici y paradas", "Detectó que EcoBici trae el barrio vacío y las paradas con erratas",
     "Punto en polígono: solo 2 estaciones y 7 paradas quedan en el borde", "Correcto", "ba_ecobici_estaciones",
     "Un campo «barrio» presente no garantiza que tenga datos."),
    (S2, "Normalizar los kilos de RAEE", "Regla de coma de miles", "Falló: la columna mezcla coma decimal y de miles; 7 valores con erratas o ambiguos",
     "Parcial", "ba_puntos_verdes", "Ante un valor ambiguo («4,033») no se adivina: queda sin dato y se cuenta."),
    (S2, "Estimar cuánto pesaría lago/raw después de comprimir", "Unos 60 MB", "Midió 108 MB: subestimó los molinetes (79 MB ya comprimidos)",
     "Incorrecto", "", "Las estimaciones se miden después y se corrigen en voz alta."),
    (S2, "Recontar los temas del lago para el README", "Escribió 23 temas", "ls lago/*.json da 22", "Incorrecto", "",
     "Las cifras de la prosa se calculan, no se escriben de memoria (regla 4 de las vistas)."),
]

# ---------------------------------------------------------------- Escribir el libro
wb = openpyxl.load_workbook(PLANTILLA)


def ampliar(ws, ultima_col):
    """Formato de la fila 5 hasta FILAS y listas desplegables extendidas."""
    for r in range(45, FILAS + 1):
        for c in range(1, ultima_col + 1):
            ws.cell(r, c)._style = copy.copy(ws.cell(5, c)._style)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(str(dv.sqref).replace("44", str(FILAS)))


def escribir(ws, filas):
    if len(filas) > FILAS - PRIMERA + 1:
        raise SystemExit(f"{ws.title}: {len(filas)} filas no caben (máx {FILAS - PRIMERA + 1})")
    for i, fila in enumerate(filas):
        for c, v in enumerate(fila, 1):
            ws.cell(PRIMERA + i, c, v)


# Hoja 1
f1 = []
for i, (k, m) in enumerate(INTEGRADAS.items(), 1):
    entidad, tipo, url, fmt, como, lic, personas, cob, periodo, frec, encontrada, obs = m
    f1.append([f"F{i:02d}", fuente_lago[k]["nombre"], entidad, tipo, url, fmt, como, lic or fuente_lago[k]["licencia"], personas,
               cob, periodo, frec, probado_de[k], "Integrada al lago", encontrada, f"id en el lago: {k}. {obs}"])
for j, (nombre, entidad, tipo, url, fmt, como, lic, personas, cob, periodo, frec, probado, estado, encontrada, obs) in enumerate(NO_ENTRARON, len(f1) + 1):
    f1.append([f"F{j:02d}", nombre, entidad, tipo, url, fmt, como, lic, personas, cob, periodo, frec, probado, estado, encontrada, obs])

hojas = {"1. Fuentes": (f1, 16), "2. Diccionario": (filas_dic, 13),
         "3. Bitácora IA": ([[i, h, enc, resp, comp, res, fid.get(f, f), apr] for i, (h, enc, resp, comp, res, f, apr) in enumerate(BITACORA, 1)], 8)}
for nombre, (filas, cols) in hojas.items():
    ws = wb[nombre]
    ampliar(ws, cols)
    escribir(ws, filas)

res = wb["Resumen"]
for fila in res.iter_rows():
    for c in fila:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.value = re.sub(r"(\$?[A-Z]+\$?)44\b", rf"\g<1>{FILAS}", c.value)
ins = wb["Instrucciones"]
ins["B5"] = "Ciudad elegida: Buenos Aires (Ciudad Autónoma de Buenos Aires, Argentina)"
ins["B7"] = "Enlace al repositorio o a la carpeta del lago: (completar) · carpeta del proyecto: lago/ · generado por entrega/generar_entrega1.py"
wb.calculation.fullCalcOnLoad = True  # openpyxl no calcula: Excel recalcula el Resumen al abrir
wb.save(SALIDA)
print(f"{SALIDA.name}: {len(f1)} fuentes ({len(INTEGRADAS)} integradas, {len(NO_ENTRARON)} fuera), "
      f"{len(filas_dic)} campos en el diccionario, {len(BITACORA)} encargos a la IA")
